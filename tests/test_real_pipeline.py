"""
Integration test for ClearSight Unified Real Pipeline.
Verifies:
1. True Crosstab on Brand_Preference by Region matches auditor findings (46.6% Total, 60.0% NCR).
2. Dual significance testing yields genuine pairwise letters and benchmark markers.
3. Top-2-Box and Scale Mean metrics compute actual sample stats.
4. Taglish qualitative coding codes actual open-ended feedback.
5. Excel Banner Book exports clean unquoted significance symbols.
"""

import os
import openpyxl
import pandas as pd
from engine.tabulation_engine import build_crosstab_table
from engine.export_engine import generate_excel_banner_book
from engine.taglish_nlp import batch_code_open_ends, scrub_pii

def test_real_pipeline():
    csv_path = os.path.join(os.path.dirname(__file__), "..", "data", "sample_survey.csv")
    df = pd.read_csv(csv_path)

    # 1. Test Tabulation on Brand_Preference by Region
    tab = build_crosstab_table(df, "Brand_Preference", ["Total", "National Capital Region (NCR)", "Balance Luzon", "Visayas", "Mindanao"])
    
    assert tab["unweighted_bases"][0] == 412
    assert tab["unweighted_bases"][1] == 120 # NCR
    assert tab["unweighted_bases"][2] == 150 # Balance Luzon

    # Brand A check (QA audit benchmark: 46.6% Total, 60.0% NCR)
    brand_a_row = next(r for r in tab["rows"] if "Brand A" in r["label"])
    assert brand_a_row["values"][0] == "46.6%", f"Expected 46.6%, got {brand_a_row['values'][0]}"
    assert brand_a_row["values"][1] == "60.0%", f"Expected 60.0%, got {brand_a_row['values'][1]}"
    assert "B" in brand_a_row["sig_letters"][1], f"Expected letter B in NCR, got {brand_a_row['sig_letters'][1]}"
    assert brand_a_row["sig_benchmarks"][1] == "++", f"Expected ++ in NCR, got {brand_a_row['sig_benchmarks'][1]}"

    # 2. Test Overall_CSAT Top-2-Box & Scale Mean
    csat_tab = build_crosstab_table(df, "Overall_CSAT", ["Total", "Region"])
    t2b_row = next(r for r in csat_tab["rows"] if r["is_net"])
    assert t2b_row["values"][0] == "85.7%", f"Expected 85.7% T2B, got {t2b_row['values'][0]}"

    mean_tab = build_crosstab_table(df, "Overall_CSAT", ["Total", "Region"], metric="mean")
    assert mean_tab["rows"][0]["values"][0] == "4.24", f"Expected 4.24, got {mean_tab['rows'][0]['values'][0]}"

    # 3. Test Real Taglish Coding
    verbatims = df["Open_Feedback"].dropna().tolist()
    nlp_res = batch_code_open_ends(verbatims)
    assert nlp_res["total_analyzed"] == 412
    assert len(nlp_res["codeframe"]) >= 3
    theme_names = [t["theme"] for t in nlp_res["codeframe"]]
    assert any("Sulit" in name or "Affordable" in name for name in theme_names)
    assert any("Expensive" in name for name in theme_names)

    # 4. Test Excel Banner Book Export
    meta = {
        "date_range": "September - October 2026",
        "unweighted_n": 412,
        "weighted_n": 412.0,
        "effective_n": 370.5,
        "efficiency_pct": 89.93
    }
    tmp_out = "/tmp/test_pipeline_export.xlsx"
    generate_excel_banner_book(tmp_out, "QA Verified Survey", [tab], meta)
    wb = openpyxl.load_workbook(tmp_out)
    ws = wb[wb.sheetnames[1]]
    
    # Check that NCR ++ is written cleanly without an apostrophe
    val_ncr_bm = ws.cell(14, 4).value
    assert val_ncr_bm == "++", f"Expected '++', got {repr(val_ncr_bm)}"

    print("[✓] ALL INTEGRATION TESTS PASSED! True end-to-end mathematical verification confirmed.")

if __name__ == "__main__":
    test_real_pipeline()
