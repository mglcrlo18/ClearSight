"""
Unit tests for unlimited columns and dynamic banner tabulation.
Verifies column letter sequencing, collision-free codes, and multi-banner expansion.
"""

from openpyxl.utils import get_column_letter
from server import ClearSightRequestHandler, load_bundled_sample, SESSION


def test_column_letter_generation():
    """Verifies that column letters sequence without overflow or collisions for unlimited columns."""
    assert get_column_letter(1) == "A"
    assert get_column_letter(2) == "B"
    assert get_column_letter(26) == "Z"
    assert get_column_letter(27) == "AA"
    assert get_column_letter(28) == "AB"
    assert get_column_letter(52) == "AZ"
    assert get_column_letter(53) == "BA"
    assert get_column_letter(702) == "ZZ"
    assert get_column_letter(703) == "AAA"


def test_tabulate_unlimited_columns():
    """Verifies that execute_tabulation generates correct multi-column tables with distinct letters and bases."""
    load_bundled_sample()
    df = SESSION.get("df")
    handler = ClearSightRequestHandler.__new__(ClearSightRequestHandler)

    banner_cols = [
        "Total",
        "NCR (A)", "Balance Luzon (B)", "Visayas (C)", "Mindanao (D)",
        "Gen Z (18–27) (E)", "Millennials (28–43) (F)", "Gen X (44–59) (G)",
        "Male (H)", "Female (I)",
        "Class ABC (J)", "Class D (K)", "Class E (L)",
        "Custom City 1 (M)", "Custom City 2 (N)", "Custom City 3 (O)"
    ]
    stubs = ["Brand Preference", "Overall CSAT"]

    tables = handler.execute_tabulation(df, banner_cols, stubs, confidence=95, fdr_enabled=True)
    assert len(tables) == 2

    t1 = tables[0]
    assert t1["title"] == "Tabulation: Brand Preference"
    assert len(t1["banner_cols"]) == 16
    assert len(t1["col_letters"]) == 16
    assert t1["col_letters"][0] == "Total"
    assert t1["col_letters"][1] == "A"
    assert t1["col_letters"][4] == "D"
    assert t1["col_letters"][5] == "E"
    assert t1["col_letters"][15] == "O"

    # Verify rows exist and have 16 values each
    assert len(t1["rows"]) > 0
    for row in t1["rows"]:
        assert len(row["values"]) == 16
        assert len(row["sig_letters"]) == 16
        assert len(row["sig_benchmarks"]) == 16

    print("PASS: test_tabulate_unlimited_columns")
