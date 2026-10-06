"""
Tests for ClearSight's Vertical Codeframe Exporter (engine/export_engine.py).
Validates:
1. Generation of vertical hierarchical .xlsx workbooks.
2. Native collapsible outline levels (outlineLevel and summaryBelow=False).
3. Pre-flight formula injection defense (CWE-1236).
4. Round-trip compatibility: export -> parse_excel_codeframe.
"""
import os
import io
import tempfile
import unittest
import openpyxl

from engine.export_engine import generate_vertical_codeframe_excel, generate_coded_hierarchy_percent_excel, build_coded_hierarchy_table
from engine.codeframe_excel_parser import parse_excel_codeframe
import server


class TestExportVerticalCodeframe(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.sample_codeframe = {
            "schema_version": 1,
            "id": "frontier_q10_test",
            "name": "Q10. Household Situation",
            "domain": "governance",
            "topics": [
                {
                    "id": "paved_roads_31",
                    "net": "GAVE FAVORABLE COMMENTS (NET)",
                    "subnet": "Infrastructure (Subnet) > Infrastructure in General (Sub-Subnet)",
                    "codes": {
                        "pos": {"code_id": 31, "label": "Roads have been paved/repaired"}
                    },
                    "keywords": ["paved", "repaired", "kalsada"],
                    "exemplars": ["Naaayos ang mga kalsada (sementado)"],
                    "dp_instruction": "Lump to code 001"
                },
                {
                    "id": "widened_roads_32",
                    "net": "GAVE FAVORABLE COMMENTS (NET)",
                    "subnet": "Infrastructure (Subnet) > Infrastructure in General (Sub-Subnet)",
                    "codes": {
                        "pos": {"code_id": 32, "label": "Roads have been widened"}
                    },
                    "keywords": ["widened", "pinalawak"],
                    "exemplars": ["Pinalawak ang mga kalsada"],
                    "dp_instruction": "Lump to code 001"
                },
                {
                    "id": "high_prices_101",
                    "net": "GAVE UNFAVORABLE COMMENTS (NET)",
                    "subnet": "Economic Hardship (Subnet)",
                    "codes": {
                        "neg": {"code_id": 101, "label": "=1+1 Malicious Formula in Label"}
                    },
                    "keywords": ["prices", "mahal"],
                    "exemplars": ["@evil_macro_command in verbatim"],
                    "dp_instruction": "+2-3 malicious formula"
                }
            ]
        }

    def test_export_creates_valid_excel_with_outlines(self):
        out_path = os.path.join(self.temp_dir, "vertical_codeframe_test.xlsx")
        res = generate_vertical_codeframe_excel(
            out_path,
            self.sample_codeframe,
            project_title="Frontier 2022 Palawan Study",
            question_text="Q10. Sa kabuuan, ano ang sitwasyon sa tahanan?"
        )
        self.assertTrue(os.path.exists(res))

        wb = openpyxl.load_workbook(out_path)
        ws = wb["Vertical Codeframe"]

        # Check outline property
        self.assertFalse(ws.sheet_properties.outlinePr.summaryBelow)

        # Check headers
        self.assertEqual(ws.cell(row=4, column=1).value, "Codes")
        self.assertEqual(ws.cell(row=4, column=2).value, "Theme / Standardized Label")
        self.assertEqual(ws.cell(row=4, column=3).value, "Anchored Verbatims (Raw Quotes)")
        self.assertEqual(ws.cell(row=4, column=4).value, "DP / Coding Instructions")

        # Check that codes are present
        found_31 = False
        found_101 = False
        found_sanitized_label = False
        found_sanitized_verbatim = False

        for r in range(5, ws.max_row + 1):
            val_a = ws.cell(row=r, column=1).value
            val_b = ws.cell(row=r, column=2).value
            val_c = ws.cell(row=r, column=3).value
            val_d = ws.cell(row=r, column=4).value

            if val_a == 31:
                found_31 = True
                self.assertIn("Roads have been paved/repaired", str(val_b))
                self.assertIn("Naaayos ang mga kalsada", str(val_c))
                self.assertEqual(val_d, "Lump to code 001")
                # Sub-sub-subnet leaf outline level should be 3
                self.assertEqual(ws.row_dimensions[r].outlineLevel, 3)

            if val_a == 101:
                found_101 = True
                # CWE-1236 check: must start with apostrophe
                if "'=1+1" in str(val_b):
                    found_sanitized_label = True
                if "'@evil_macro" in str(val_c):
                    found_sanitized_verbatim = True

        self.assertTrue(found_31)
        self.assertTrue(found_101)
        self.assertTrue(found_sanitized_label)
        self.assertTrue(found_sanitized_verbatim)

    def test_round_trip_parse_after_export(self):
        """Export vertical codeframe to Excel, then parse it back with codeframe_excel_parser."""
        out_path = os.path.join(self.temp_dir, "round_trip_test.xlsx")
        generate_vertical_codeframe_excel(
            out_path,
            self.sample_codeframe,
            project_title="Frontier Palawan",
            question_text="Q10. Household Situation"
        )

        parsed_cf = parse_excel_codeframe(out_path, codeframe_name="Round Trip Tested Codeframe")
        self.assertIn("topics", parsed_cf)
        codes_found = []
        for t in parsed_cf["topics"]:
            for pol, c_info in t["codes"].items():
                codes_found.append(c_info["code_id"])

        self.assertIn(31, codes_found)
        self.assertIn(32, codes_found)
        self.assertIn(101, codes_found)

    def test_generate_coded_hierarchy_percent_excel(self):
        """Verify generation of coded vertical hierarchy table with percentages matching screenshot."""
        records = [
            {"response_id": 1, "assigned_codes": [31], "assigned_themes": ["Roads have been paved/repaired"]},
            {"response_id": 2, "assigned_codes": [31], "assigned_themes": ["Roads have been paved/repaired"]},
            {"response_id": 3, "assigned_codes": [32], "assigned_themes": ["Roads have been widened"]},
            {"response_id": 4, "assigned_codes": [101], "assigned_themes": ["=1+1 Malicious Formula in Label"]}
        ]
        total_n = 100

        rows = build_coded_hierarchy_table(records, self.sample_codeframe, total_base=total_n)
        self.assertTrue(len(rows) > 0)

        # Check types
        types = [r["type"] for r in rows]
        self.assertIn("net", types)
        self.assertIn("subnet", types)
        self.assertIn("leaf", types)

        # Verify deduplicated reach:
        # Favorable NET should contain resps 1, 2, 3 -> 3 unique resps / 100 = 3%
        fav_net = next(r for r in rows if "FAVORABLE" in r["label"])
        self.assertEqual(fav_net["pct_str"], "3")

        # Leaf 31 has 2 resps / 100 = 2%
        leaf_31 = next(r for r in rows if r.get("code_id") == 31)
        self.assertEqual(leaf_31["pct_str"], "2")

        # Generate Excel
        out_path = os.path.join(self.temp_dir, "coded_hierarchy_test.xlsx")
        generate_coded_hierarchy_percent_excel(
            out_path,
            rows,
            project_title="Frontier Palawan Study",
            question_text="Q10. Household Situation",
            total_n=total_n
        )
        self.assertTrue(os.path.exists(out_path))

        wb = openpyxl.load_workbook(out_path)
        ws = wb["Coded Hierarchy (%)"]

        # Check headers (Row 5)
        self.assertEqual(ws.cell(row=5, column=1).value, "Theme / Standardized Response Hierarchy")
        self.assertEqual(ws.cell(row=5, column=2).value, "%")

        # Check pale yellow header fill on Col B
        self.assertIn("FEF9C3", str(ws.cell(row=5, column=2).fill.start_color.rgb))

        # Check data row styling and values
        cell_fav_pct = None
        for r in range(6, 6 + len(rows)):
            lbl = ws.cell(row=r, column=1).value
            pct = ws.cell(row=r, column=2).value
            fill_b = ws.cell(row=r, column=2).fill
            # Pale yellow background on Col B data cells
            self.assertIn("FFFDE7", str(fill_b.start_color.rgb))
            if lbl and "GAVE FAVORABLE" in str(lbl):
                cell_fav_pct = pct

        self.assertEqual(cell_fav_pct, 3)

    def test_server_export_coded_hierarchy_endpoint(self):
        """Verify /api/export/coded-hierarchy-excel streams valid Excel workbook."""
        server.load_bundled_sample()
        with server.SESSION_LOCK:
            server.SESSION["open_feedback_analysis"] = {
                "total_analyzed": 50,
                "records": [
                    {"response_id": 1, "assigned_themes": ["Affordable / High Value (Sulit)"], "assigned_codes": [110]},
                    {"response_id": 2, "assigned_themes": ["Expensive / High Pricing Friction"], "assigned_codes": [120]}
                ]
            }

        handler = server.ClearSightRequestHandler.__new__(server.ClearSightRequestHandler)
        captured_headers = []
        handler.send_response = lambda status: captured_headers.append(("status", status))
        handler.send_header = lambda k, v: captured_headers.append((k, v))
        handler.end_headers = lambda: None
        handler.wfile = io.BytesIO()

        handler.handle_export_coded_hierarchy_excel()
        excel_bytes = handler.wfile.getvalue()
        self.assertGreater(len(excel_bytes), 1000)

        # Verify workbook opens cleanly
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Coded Hierarchy (%)", wb.sheetnames)
        ws = wb["Coded Hierarchy (%)"]
        self.assertIn("%", str(ws.cell(row=5, column=2).value))


if __name__ == "__main__":
    unittest.main()
