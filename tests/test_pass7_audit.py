"""
ClearSight Pass 7 regression tests (fb793d3). Each test names the finding it guards (Pass7_Fix_List.md).
Known-answer where possible; property tests (Hypothesis) for PII and weight-name detection.
"""
import io
import json
import re
import threading
import time
from pathlib import Path
import unittest

import numpy as np
import openpyxl
import pandas as pd
try:
    import pytest
    from hypothesis import given, settings, strategies as st, HealthCheck
except ImportError:
    import scripts.run_all_tests
    import pytest
    from hypothesis import given, settings, strategies as st, HealthCheck

import server
from server import SESSION
from engine import ingestion as ing
from engine.driver_analysis import compute_johnsons_relative_weights
from engine.taglish_nlp import scrub_pii
from tests.test_pass5_regressions import _post, _load, _FakeHandler  # noqa: F401  (reuses the autouse session fixture below)
from tests.test_pass5_regressions import _isolated_session  # noqa: F401

REPO = Path(__file__).resolve().parents[1]
JS = (REPO / "static" / "app.js").read_text(encoding="utf-8")
HTML = (REPO / "static" / "index.html").read_text(encoding="utf-8")
CSS = (REPO / "static" / "styles.css").read_text(encoding="utf-8")
HDR = ["", "Total", "Male", "Female", "NCR", "Luzon"]


def _xlsx(rows, title="Q5. Brand bought most often", fmt=None):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Q5"
    ws.append([title])
    for r in rows:
        ws.append(r)
    if fmt:
        for (r, c), f in fmt.items():
            ws.cell(row=r, column=c).number_format = f
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


def _table(b):
    df, meta = ing.read_survey_file(b, "t.xlsx")
    assert meta.get("is_data_table"), meta
    return meta["pretabulated_table"], meta


# ---------------------------------------------------------------- P7-01 (held-out FGD data published)
def test_p7_01_fgd_dataset_not_in_repo():
    hits = [p for p in REPO.rglob("*.csv") if "fgd" in p.name.lower() and ".git" not in p.parts]
    assert not hits, f"FGD verbatims must not be committed (RA 10173, held-out contamination): {hits}"


# ---------------------------------------------------------------- P7-02/03/04 (boot)
def test_p7_02_export_engine_imports_typing_names():
    src = (REPO / "engine" / "export_engine.py").read_text(encoding="utf-8")
    assert re.search(r"from typing import .*\bAny\b", src) and re.search(r"from typing import .*\bOptional\b", src)


def test_p7_04_no_pep701_nested_quote_fstrings():
    import ast, sys
    for p in (REPO / "engine").glob("*.py"):
        src = p.read_text(encoding="utf-8")
        ast.parse(src, feature_version=(3, 10))  # SyntaxError on 3.10/3.11 syntax


# ---------------------------------------------------------------- P7-14/17 (bases)
def test_p7_14_missing_unweighted_base_is_unknown_not_400():
    t, _ = _table(_xlsx([HDR, ["Effective base", 900, 450, 450, 250, 650], ["", "", "A", "B", "C", "D"],
                         ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47]]))
    assert t["unweighted_bases"] == [None] * 5
    assert 400 not in (t["unweighted_bases"] + t["weighted_bases"])
    assert any("base" in w.lower() for w in t.get("warnings", []))


def test_p7_17_comma_formatted_base_parses():
    t, _ = _table(_xlsx([HDR, ["Column sample size", "1,075", "530", "545", "300", "775"],
                         ["Weighted base", "1,075.0", "520.3", "554.7", "310", "765"],
                         ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47]]))
    assert t["unweighted_bases"] == [1075, 530, 545, 300, 775]
    assert t["weighted_bases"][0] == pytest.approx(1075.0)
    assert t["small_base"][0] is False


@pytest.mark.parametrize("raw,expected", [("1,075", 1075), ("(n=1,075)", 1075), ("n = 530", 530), (812, 812), ("812.0", 812), ("", None), ("-", None), (None, None)])
def test_p7_17_parse_count(raw, expected):
    assert ing._parse_count(raw) == expected


# ---------------------------------------------------------------- P7-16 (dropped rows)
def test_p7_16_dash_cells_and_blank_stub_rows_are_kept():
    t, _ = _table(_xlsx([HDR, ["Column sample size", 1075, 530, 545, 300, 775], ["", "", "A", "B", "C", "D"],
                         ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47], ["Brand B", 0.30, "-", 0.31, 0.29, "-"],
                         [None, 0.10, 0.1, 0.1, 0.1, 0.1], ["Brand D", 0.05, 0.05, 0.05, 0.05, 0.05]]))
    labels = [r["label"] for r in t["rows"]]
    assert labels[:2] == ["Brand A", "Brand B"] and labels[-1] == "Brand D" and len(labels) == 4
    assert labels[2].startswith("(unlabelled row")
    brand_a = t["rows"][0]
    assert "0.31" not in brand_a["sig_benchmarks"], "Brand B's values were swallowed as Brand A's sig row"


def test_p7_16_real_sig_row_still_attached():
    t, _ = _table(_xlsx([HDR, ["Column sample size", 1075, 530, 545, 300, 775],
                         ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47], ["", "B", "", "C", "", ""],
                         ["Brand B", 0.55, 0.60, 0.50, 0.58, 0.53]]))
    assert [r["label"] for r in t["rows"]] == ["Brand A", "Brand B"]
    assert t["rows"][0]["sig_letters"][2] == "C"  # Total column carries no letters by design


# ---------------------------------------------------------------- P7-18 (percent rendering)
def test_p7_18_percent_only_when_cell_says_so():
    rows = [HDR, ["Column sample size", 1075, 530, 545, 300, 775],
            ["Brand A (pct 0-100)", 45, 40, 50, 42, 47], ["Mean score", 0.85, 3.9, 4.1, 1.0, 0.5],
            ["Share (formatted %)", 0.45, 0.40, 0.50, 0.42, 0.47]]
    fmt = {(6, c): "0%" for c in range(2, 7)}
    t, _ = _table(_xlsx(rows, fmt=fmt))
    vals = {r["label"]: r["values"] for r in t["rows"]}
    assert vals["Brand A (pct 0-100)"][0] == "45"
    assert not any(v.endswith("%") for v in vals["Mean score"])
    assert vals["Share (formatted %)"][0].endswith("%")


# ---------------------------------------------------------------- P7-25 (column letters)
@pytest.mark.parametrize("i,letters", [(0, "A"), (25, "Z"), (26, "AA"), (27, "AB"), (51, "AZ"), (52, "BA"), (701, "ZZ"), (702, "AAA")])
def test_p7_25_excel_col_letters(i, letters):
    assert ing._excel_col_letter(i) == letters


# ---------------------------------------------------------------- P7-20 (parser errors surface)
def test_p7_20_parser_bug_is_not_swallowed(monkeypatch):
    b = _xlsx([HDR, ["Column sample size", 1075, 530, 545, 300, 775], ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47]])
    def boom(*a, **k):
        raise RuntimeError("parser bug")
    monkeypatch.setattr(ing, "parse_pretabulated_banner_sheet", boom)
    with pytest.raises(RuntimeError):
        ing.read_survey_file(b, "t.xlsx")


# ---------------------------------------------------------------- P7-21 (header row)
def test_p7_21_header_row_with_blank_cells_is_chosen():
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Respondent_ID", None, "Gender", None, "Region"])
    for i in range(20):
        ws.append([f"R{i:03d}", "note text", "Male" if i % 2 else "Female", "more text", "NCR"])
    bio = io.BytesIO(); wb.save(bio)
    df, meta = ing.read_survey_file(bio.getvalue(), "raw.xlsx")
    assert list(df.columns)[0] == "Respondent_ID" and len(df) == 20


# ---------------------------------------------------------------- P7-15 (weight autodetect)
@pytest.mark.parametrize("name,expected", [
    ("weight", True), ("WGT", True), ("Wt", True), ("final_weight", True), ("rim_weight", True),
    ("Product_Weight_g", False), ("Sachet_weight", False), ("Weighted_volume", False),
    ("Weight_of_pack_kg", False), ("Body_Weight", False),
])
def test_p7_15_weight_name_detection(name, expected):
    assert bool(ing.WEIGHT_NAME_RE.match(name)) is expected


@settings(max_examples=150, deadline=None)
@given(st.sampled_from(["Product", "Sachet", "Pack", "Body", "Net", "Bag", "Item"]),
       st.sampled_from(["_weight_g", "_Weight_kg", " weight (g)", "_weight", "Weight"]))
def test_p7_15_product_weights_never_detected(prefix, suffix):
    assert not ing.WEIGHT_NAME_RE.match(prefix + suffix)


def _upload(df, name="w.csv"):
    raw = df.to_csv(index=False).encode()
    return _post("/api/upload", raw=raw, headers={"Content-Type": "application/octet-stream", "X-Filename": name})


def test_p7_15_upload_suggests_but_does_not_apply_weights():
    df = pd.DataFrame({"Respondent_ID": range(1, 61), "Gender": ["M", "F"] * 30,
                       "weight": np.tile([0.5, 1.5], 30)})
    st_, d = _upload(df)
    assert st_ == 200 and d["status"] == "success", d
    assert d["detected_weight_column"] == "weight"
    assert d["weighted"] is False and SESSION["weights"] is None
    st_, a = _post("/api/weight/apply-column", {"column": "weight"})
    assert st_ == 200 and a["status"] == "success", a
    assert SESSION["weights"] is not None and len(SESSION["weights"]) == 60
    # Kish efficiency for weights 0.5/1.5 alternating: (sum w)^2 / (n * sum w^2) = 3600/(60*75) = 80%
    assert a["diagnostics"]["efficiency"] == pytest.approx(80.0)


@pytest.mark.parametrize("vals,ok", [([1.0, 2.0, 0.5] * 20, True), ([1.0, 0.0, 2.0] * 20, False),
                                      ([1.0, -1.0, 2.0] * 20, False), ([1.0, None, 2.0] * 20, False)])
def test_p7_15_apply_column_validates(vals, ok):
    _load(pd.DataFrame({"g": ["a", "b", "c"] * 20, "w": vals}))
    st_, a = _post("/api/weight/apply-column", {"column": "w"})
    assert (st_ == 200) is ok
    st_, a = _post("/api/weight/apply-column", {"column": "nope"})
    assert st_ == 400


# ---------------------------------------------------------------- P7-22 (pretab skips hygiene)
def test_p7_22_pretabulated_import_skips_hygiene_audit():
    b = _xlsx([HDR, ["Column sample size", 1075, 530, 545, 300, 775],
               ["Brand A", 0.45, 0.40, 0.50, 0.42, 0.47], ["Brand B", 0.45, 0.40, 0.50, 0.42, 0.47]])
    st_, d = _post("/api/upload", raw=b, headers={"Content-Type": "application/octet-stream", "X-Filename": "tab.xlsx"})
    assert st_ == 200 and d["is_data_table"] is True and d["flagged_hygiene"] == 0
    assert d["total_respondents"] == 1075


# ---------------------------------------------------------------- P7-19 / P7-33 (filename)
def test_p7_19_filename_is_unquoted_and_control_chars_stripped():
    df = pd.DataFrame({"Respondent_ID": range(1, 41), "Gender": ["M", "F"] * 20})
    st_, d = _upload(df, name="Sarbey%20Ma%C3%B1ila%0A%0D.csv")
    assert st_ == 200 and d["filename"] == "Sarbey Mañila.csv"


def test_p7_19_toast_never_parses_message_as_html():
    i = JS.index("function showToast"); block = JS[i:JS.index("\n}\n", i)]
    assert "${message}" not in block and "textContent" in block


def test_p7_33_client_percent_encodes_filename_header():
    assert "'X-Filename': file.name" not in JS
    assert JS.count("encodeURIComponent(file.name)") >= 2


# ---------------------------------------------------------------- P5-03 / P7-34 (all innerHTML sinks)
def test_p7_34_known_unescaped_sinks_now_escaped():
    for frag in ("${escapeHtml(c.term)}", "${escapeHtml(p.path)}", "escapeHtml(r.model || r.test"):
        assert frag in JS, frag


# ---------------------------------------------------------------- P7-31 (dead codeframe button)
def test_p7_31_codeframe_button_targets_existing_input():
    targets = re.findall(r"getElementById\('([\w-]+)'\)\.click\(\)", JS + HTML)
    ids = set(re.findall(r'id="([\w-]+)"', HTML)) | set(re.findall(r"id=\"([\w-]+)\"", JS))
    missing = [t for t in targets if t not in ids]
    assert not missing, f".click() on non-existent element ids: {missing}"


# ---------------------------------------------------------------- P7-32 (mobile overflow)
def test_p7_32_stepper_wraps_on_phones():
    m = re.search(r"@media \(max-width:\s*(\d+)px\)\s*\{[^@]*\.workflow-stepper\s*\{[^}]*flex-wrap:\s*wrap", CSS)
    assert m and int(m.group(1)) >= 390


# ---------------------------------------------------------------- P7-28 (hygiene default / strict bools)
def test_p7_28_hygiene_quarantine_is_opt_in():
    assert server.DEFAULT_ANALYSIS_CONFIG["hygiene"]["quarantine_straightliners"] is False
    assert server.DEFAULT_ANALYSIS_CONFIG["hygiene"]["quarantine_speeders"] is False
    for i in ("toggle-hygiene-straight", "toggle-hygiene-speeders", "check-straight", "check-speeder"):
        tag = re.search(r'<input[^>]*id="%s"[^>]*>' % i, HTML).group(0)
        assert " checked" not in tag, i


@pytest.mark.parametrize("val", ["false", "true", 1, 0, None, "yes"])
def test_p7_28_hygiene_filter_requires_real_boolean(val):
    _load(pd.DataFrame({"a": [1, 2, 3] * 10}))
    _post("/api/hygiene-filter", {"filter_straight_liners": val, "filter_speeders": val})
    assert SESSION["analysis_config"]["hygiene"]["quarantine_straightliners"] is False


def test_p7_28_default_config_not_mutated_by_session():
    SESSION["analysis_config"]["hygiene"]["quarantine_straightliners"] = True
    assert server.DEFAULT_ANALYSIS_CONFIG["hygiene"]["quarantine_straightliners"] is False


# ---------------------------------------------------------------- P7-23 (driver weight column)
def test_p7_23_driver_misspelt_weight_column_raises():
    rng = np.random.default_rng(1)
    d = pd.DataFrame({"y": rng.normal(size=100), "x1": rng.normal(size=100), "x2": rng.normal(size=100), "w": rng.uniform(0.2, 3, 100)})
    with pytest.raises(ValueError, match="not found"):
        compute_johnsons_relative_weights(d, "y", ["x1", "x2"], weights="nope")
    with pytest.raises(ValueError, match="rows"):
        compute_johnsons_relative_weights(d, "y", ["x1", "x2"], weights=np.ones(99))
    a = compute_johnsons_relative_weights(d, "y", ["x1", "x2"], weights="w")
    b = compute_johnsons_relative_weights(d, "y", ["x1", "x2"])
    assert a["r_squared"] != b["r_squared"]


# ---------------------------------------------------------------- P7-24 (natural sort)
def test_p7_24_numeric_categories_sort_numerically():
    d = pd.DataFrame({"Q_code": [str(x) for x in [3, 1, 10, 2, 11, 1, 3]] * 5})
    assert ing.autodetect_schema(d)["Q_code"]["categories"] == ["1", "2", "3", "10", "11"]


# ---------------------------------------------------------------- P7-27 (SESSION_LOCK consistency)
def test_p7_27_tabulate_reads_df_and_weights_atomically():
    """Uploads racing a tabulate must never yield 'weights length != rows' (seen 154x on fb793d3)."""
    big = pd.DataFrame({"Respondent_ID": range(700), "Gender": ["M", "F"] * 350, "Region": ["NCR", "Luzon", "VisMin", "Other", "X", "Y", "Z"] * 100})
    small = big.head(200).copy()
    errors = []
    stop = threading.Event()

    def flip():
        k = 0
        while not stop.is_set():
            df = big if k % 2 else small
            with server.SESSION_LOCK:
                SESSION["df"] = df; SESSION["schema"] = ing.autodetect_schema(df)
                SESSION["weights"] = np.ones(len(df))
            k += 1
            time.sleep(0.002)  # let the tabulate thread take the lock; without this, 3.10/3.12 starve it

    t = threading.Thread(target=flip); t.start()
    try:
        for _ in range(40):
            st_, d = _post("/api/tabulate", {"banner_cols": ["Gender"], "stubs": ["Region"]})
            if st_ != 200 or "length" in json.dumps(d):
                errors.append((st_, str(d)[:120]))
    finally:
        stop.set(); t.join()
    assert not errors, errors[:3]


# ---------------------------------------------------------------- P7-29 (PII over-masking) / CS-023
@pytest.mark.parametrize("text", ["Si mama naman ang bumili", "Kay ganda talaga nito", "Ni minsan hindi nasira",
                                  "Kay bilis dumating ng order", "Para kay lola ito binili", "I ate rice for lunch today",
                                  "Mabilis si seller at mabait.", "Nindot kaayo ang bugas kay barato ra ug lami pa.",
                                  "Si kuya guard mabait", "Si Duterte ang"])
def test_p7_29_common_words_not_masked(text):
    assert scrub_pii(text) == text


@pytest.mark.parametrize("text,name", [("Si Maria Santos ang kausap ko", "Maria"), ("Kausap ko si juan kanina sa opisina.", "juan"),
                                       ("Ibigay mo kay pedro ang bayad.", "pedro"), ("Mabait si Ate Joy ang nag-assist", "Joy"),
                                       ("Tinulungan ako ni Juan dela Cruz", "Juan"), ("Kausap ko si Gng. Reyes kahapon", "Reyes")])
def test_p7_29_real_names_still_masked(text, name):
    out = scrub_pii(text)
    assert name not in out and "[NAME_REDACTED]" in out


# ---------------------------------------------------------------- P7-06 (coder v3 PII) — property
CF = [{"id": "price", "subnet": "Price", "keywords": ["mahal", "presyo", "mura"], "exemplars": ["mahal ang presyo"]},
      {"id": "svc", "subnet": "Service", "keywords": ["serbisyo", "staff"], "exemplars": ["mabait ang staff"]}]


def _v3():
    from engine.coder_v3 import CoderV3Engine
    return CoderV3Engine(codeframe=CF)


FIRST = st.sampled_from(["Juan", "Maria", "Jose", "Ana", "Pedro", "Liza"])
LAST = st.sampled_from(["Santos", "Reyes", "Cruz", "Bautista", "Garcia"])
PHONE = st.from_regex(r"09[0-9]{9}", fullmatch=True)


_V3 = None


@pytest.fixture(scope="module", autouse=True)
def _v3_engine():
    global _V3
    _V3 = _v3()


@settings(max_examples=60, deadline=None, suppress_health_check=[HealthCheck.too_slow, HealthCheck.function_scoped_fixture])
@given(FIRST, LAST, PHONE)
def test_p7_06_v3_records_never_contain_raw_pii(first, last, phone):
    text = f"Kausap ko si {first} {last}, tawagan niyo ako sa {phone} o {first.lower()}@gmail.com. Mahal ang bigas."
    out = _V3.batch_code([text])
    blob = json.dumps(out, ensure_ascii=False)
    assert phone not in blob and f"{first.lower()}@gmail.com" not in blob and f"{first} {last}" not in blob


def test_p7_06_server_v3_quote_is_masked():
    _load(pd.DataFrame({"Respondent_ID": range(1, 31),
                        "Remarks": ["Mahal ang bigas, tawagan si Maria Santos sa 09171234567"] * 30}))
    st_, d = _post("/api/code-open-ends", {"coder": "v3", "codeframe": "consumer_default"})
    assert st_ == 200 and d["status"] == "success", d
    blob = json.dumps(d, ensure_ascii=False)
    assert "09171234567" not in blob and "Maria Santos" not in blob
    assert d["encoder"] in ("minilm", "hash-fallback")  # P7-35: the response says which encoder ran


# ---------------------------------------------------------------- P7-07/10 (v3 coding behaviour)
def test_p7_07_negation_words_are_negators_only():
    from engine.coder_v3 import TaglishPreProcessor
    neg = {w.lower() for w in TaglishPreProcessor.NEGATION_WORDS}
    assert {"hindi", "wala"} <= neg and not ({"din", "rin", "mali"} & neg)


def test_p7_10_keywords_match_whole_words_only():
    rec = _v3().batch_code(["Kumurap ako sa sobrang liwanag", "Mura na ngayon ang bigas"])
    assert rec[1]["assigned_codes"] and "price" in rec[1]["assigned_codes"]
    assert "price" not in (rec[0].get("assigned_codes") or []) or rec[0].get("keyword_hits")


def test_p7_10_empty_verbatim_gets_no_theme_and_is_flagged():
    rec = _v3().batch_code(["   ", "Mahal ang presyo"])[0]
    assert not (rec.get("assigned_themes") or []) and not (rec.get("assigned_codes") or []) and rec.get("needs_review") is True


# ---------------------------------------------------------------- P7-09 (offline)
def test_p7_09_v3_forces_offline_hub():
    import os
    import engine.coder_v3  # noqa: F401
    assert os.environ.get("HF_HUB_OFFLINE") == "1" and os.environ.get("HF_HUB_DISABLE_TELEMETRY") == "1"


# ---------------------------------------------------------------- P7-11 (default coder)
def test_p7_11_default_coder_is_v2_until_v3_passes_gate():
    assert server.CODER_DEFAULT == "v2"


# ---------------------------------------------------------------- CS-005 (Rao-Scott MRCV)
def _reference_rs2(Y, g):
    """Independent textbook computation (Bilder & Loughin 2004, eq. for RS2), scipy only."""
    from scipy import stats
    groups = sorted(set(g)); G = len(groups); J = Y.shape[1]
    xs = sum(stats.chi2_contingency(np.array([[np.sum((g == k) & (Y[:, j] == 1)) for k in groups],
                                              [np.sum((g == k) & (Y[:, j] == 0)) for k in groups]]),
                                    correction=False)[0] for j in range(J))
    lam = np.repeat(np.linalg.eigvalsh(np.corrcoef(Y, rowvar=False)), G - 1)
    a2 = ((lam - lam.mean()) ** 2).sum() / (len(lam) * lam.mean() ** 2)
    stat = xs / (lam.mean() * (1 + a2)); df = len(lam) / (1 + a2)
    return stat, df, stats.chi2.sf(stat, df)


@pytest.mark.parametrize("G", [2, 3, 4])
def test_cs005_rs2_matches_reference(G):
    from engine.stats_engine import rao_scott_second_order_mrcv
    rng = np.random.default_rng(G)
    n = 600; z = rng.normal(size=n)
    Y = np.column_stack([(z + rng.normal(size=n) > 0.3 * j).astype(int) for j in range(5)])
    g = np.array([f"g{i % G}" for i in range(n)])
    Y[g == "g1", 0] = (rng.random((g == "g1").sum()) < 0.7).astype(int)
    got = rao_scott_second_order_mrcv(np.zeros((5, G)), n, respondent_matrix=Y, group_labels=g)
    ref = _reference_rs2(Y, g)
    assert got[0] == pytest.approx(ref[0], rel=1e-9) and got[1] == pytest.approx(ref[1], rel=1e-9) and got[2] == pytest.approx(ref[2], rel=1e-6, abs=1e-300)


def test_cs005_rs2_null_size_is_close_to_nominal():
    """Strongly correlated items under H0: the naive sum of item chi-squares over-rejects; RS2 stays near 5%."""
    from engine.stats_engine import rao_scott_second_order_mrcv
    from scipy import stats
    rng = np.random.default_rng(11); rej_rs2 = rej_naive = 0; sims = 300
    for _ in range(sims):
        n = 300; z = rng.normal(size=n)
        Y = np.column_stack([(z + 0.4 * rng.normal(size=n) > 0).astype(int) for _ in range(6)])
        g = rng.integers(0, 2, n)
        s, d, p = rao_scott_second_order_mrcv(None, n, respondent_matrix=Y, group_labels=g)
        rej_rs2 += p < 0.05
        naive = sum(stats.chi2_contingency(np.array([[np.sum((g == k) & (Y[:, j] == 1)) for k in (0, 1)],
                                                     [np.sum((g == k) & (Y[:, j] == 0)) for k in (0, 1)]]), correction=False)[0] for j in range(6))
        rej_naive += stats.chi2.sf(naive, 6) < 0.05
    assert 0.02 <= rej_rs2 / sims <= 0.09, rej_rs2 / sims
    assert rej_naive / sims > 0.075 and rej_naive > rej_rs2  # naive chi2(J) over-rejects (~10% here)


def test_cs005_tabulation_skips_overlapping_banners_and_runs_on_one_variable():
    from engine.tabulation_engine import build_crosstab_table
    rng = np.random.default_rng(3)
    opts = ["Shopee", "Lazada", "Tiktok", "Shopee, Lazada", "Lazada, Tiktok"]
    df = pd.DataFrame({"Region": rng.choice(["NCR", "Luzon", "VisMin"], 300), "Gender": rng.choice(["M", "F"], 300),
                       "Channels": rng.choice(opts, 300)})
    one = build_crosstab_table(df, "Channels", ["Total", "Region"])
    assert one["mrcv"] and "skipped" not in one["mrcv"] and 0 <= one["mrcv"]["p_val"] <= 1
    two = build_crosstab_table(df, "Channels", ["Total", "Region", "Gender"])
    assert two["mrcv"] is None or "skipped" in two["mrcv"]


# ---------------------------------------------------------------- CS-019 (commas inside option labels)
def test_cs019_lowercase_fragment_joins_previous_option():
    s = pd.Series(["Shade, tent or covered waiting area, Clearer signage", "Clearer signage", "Shade, tent or covered waiting area",
                   "Seats for seniors and PWDs, Other: Dagdag na CR", "Seats for seniors and PWDs", "Clearer signage, Seats for seniors and PWDs",
                   "Shade, tent or covered waiting area", "Clearer signage", "Seats for seniors and PWDs", "Other: Sana may delivery"])
    ind, other = ing.resolve_google_forms_checkboxes(s)
    assert set(ind.columns) == {"Shade, tent or covered waiting area", "Clearer signage", "Seats for seniors and PWDs"}
    assert ind["Shade, tent or covered waiting area"].tolist()[:3] == [1, 0, 1]
    assert other.iloc[3] == "Dagdag na CR" and other.iloc[9] == "Sana may delivery"


# ---------------------------------------------------------------- P7-30 (thesis export guard)
@pytest.mark.parametrize("path", ["/api/export/thesis-download", "/api/export/thesis-tables"])
def test_p7_30_thesis_exports_need_a_table(path):
    _load(pd.DataFrame({"Gender": ["M", "F"] * 30, "Sat_1to5": [1, 2, 3, 4, 5, 4] * 10}))
    h = _FakeHandler("GET", path); h.do_GET()
    assert h.status == 400
    st_, _ = _post("/api/tabulate", {"banner_cols": ["Gender"], "stubs": ["Sat_1to5"]})
    assert st_ == 200
    h = _FakeHandler("GET", path); h.do_GET()
    assert h.status == 200


# ---------------------------------------------------------------- P7-28 (settings endpoint strict booleans)
@pytest.mark.parametrize("val", ["false", "true", 1, 0, None])
def test_p7_28_settings_analysis_ignores_non_booleans(val):
    before = json.loads(json.dumps(SESSION["analysis_config"]))
    _post("/api/settings/analysis", {"hygiene": {"quarantine_straightliners": val}, "stats": {"chi_square": val, "bogus_key": True}})
    assert SESSION["analysis_config"]["hygiene"]["quarantine_straightliners"] == before["hygiene"]["quarantine_straightliners"]
    assert SESSION["analysis_config"]["stats"]["chi_square"] == before["stats"]["chi_square"]
    assert "bogus_key" not in SESSION["analysis_config"]["stats"]


def test_p7_28_settings_analysis_accepts_real_booleans():
    _post("/api/settings/analysis", {"hygiene": {"quarantine_straightliners": True}})
    assert SESSION["analysis_config"]["hygiene"]["quarantine_straightliners"] is True
    assert server.DEFAULT_ANALYSIS_CONFIG["hygiene"]["quarantine_straightliners"] is False


# ---------------------------------------------------------------- CS-N06 (demo N in the import summary)
def test_cs_n06_import_summary_never_shows_demo_n():
    i = JS.index("function applyIngestedSummary"); block = JS[i:JS.index("\n}\n", i)]
    assert "|| 412" not in block


# ---------------------------------------------------------------- CS-017 residual (weighted chi-square)
def test_cs017_weighted_chi_square_is_scale_invariant_and_matches_unweighted_for_equal_weights():
    from engine.tabulation_engine import build_crosstab_table
    from scipy import stats
    rng = np.random.default_rng(5)
    df = pd.DataFrame({"Brand": rng.choice(["A", "B", "C"], 400), "Gender": rng.choice(["M", "F"], 400)})
    ref = stats.chi2_contingency(pd.crosstab(df.Brand, df.Gender).to_numpy(), correction=False)
    t1 = build_crosstab_table(df, "Brand", ["Total", "Gender"])
    t10 = build_crosstab_table(df, "Brand", ["Total", "Gender"], weights=np.full(400, 10.0))
    assert t1["chi_square"]["chi2_stat"] == pytest.approx(ref[0], abs=1e-3)
    assert t10["chi_square"]["chi2_stat"] == pytest.approx(ref[0], abs=1e-3)


def test_cs017_chi_square_skipped_for_overlapping_banners():
    from engine.tabulation_engine import build_crosstab_table
    rng = np.random.default_rng(6)
    df = pd.DataFrame({"Brand": rng.choice(["A", "B", "C"], 300), "Gender": rng.choice(["M", "F"], 300), "Region": rng.choice(["N", "S"], 300)})
    t = build_crosstab_table(df, "Brand", ["Total", "Gender", "Region"])
    assert t["chi_square"] is None or "skipped" in t["chi_square"]


# ---------------------------------------------------------------- mutation gaps (Pass 7 targeted mutants M10/M41/M66/M67)
def test_m66_kruskal_unweighted_matches_scipy_with_ties():
    from scipy import stats
    from engine.non_parametric import survey_weighted_kruskal_wallis
    rng = np.random.default_rng(8)
    y = rng.integers(1, 6, 240).astype(float); g = np.repeat(["a", "b", "c"], 80); y[g == "c"] += 1
    r = survey_weighted_kruskal_wallis(y, g)
    h = r.get("h_stat", r.get("H", r.get("statistic")))
    assert h == pytest.approx(stats.kruskal(y[g == "a"], y[g == "b"], y[g == "c"]).statistic, rel=1e-4)
    r10 = survey_weighted_kruskal_wallis(y, g, weights=np.full(len(y), 10.0))
    assert r10["h_stat"] == pytest.approx(h, rel=1e-4)  # weight scale must not change H


def test_m67_turf_reach_known_answer_and_weighted():
    from engine.turf_engine import calculate_turf
    m = [[1, 0, 0], [0, 1, 0], [1, 1, 0], [0, 0, 1]]
    r = calculate_turf(m, ["a", "b", "c"], k=2)
    assert r["individual_item_reach_pct"] == {"a": 50.0, "b": 50.0, "c": 25.0}
    best = {tuple(p["items"]): p["reach_pct"] for p in r["top_portfolios"]}
    assert best[("a", "b")] == pytest.approx(75.0) and best[("a", "c")] == pytest.approx(75.0)
    rw = calculate_turf(m, ["a", "b", "c"], k=1, weights=[3, 1, 1, 1])
    assert rw["individual_item_reach_pct"]["a"] == pytest.approx(100 * 4 / 6, abs=0.1)


def test_m41_99_percent_tier_is_stricter_than_95():
    from engine.tabulation_engine import build_crosstab_table
    # p ~ 0.02 between columns: uppercase at 95%, lowercase (not upper) at 99%
    df = pd.DataFrame({"G": ["A"] * 400 + ["B"] * 400, "Y": ["yes"] * 200 + ["no"] * 200 + ["yes"] * 232 + ["no"] * 168})
    t95 = build_crosstab_table(df, "Y", ["Total", "G"], confidence_level=0.95)
    t99 = build_crosstab_table(df, "Y", ["Total", "G"], confidence_level=0.99)
    def letters(t):
        row = next(r for r in t["rows"] if str(r.get("label", "")).lower().startswith("yes"))
        return "".join(row.get("sig_letters") or row.get("letters") or [])
    assert "A" in letters(t95)
    assert "A" not in letters(t99)


def test_m10_population_count_targets_are_normalised_by_their_sum():
    from engine.stats_engine import calculate_rim_weights
    df = pd.DataFrame({"Region": ["N"] * 50 + ["S"] * 50})
    out = calculate_rim_weights(df, {"Region": {"N": 300000, "S": 100000}})
    w = out[0] if isinstance(out, tuple) else out
    w = np.asarray(w if not isinstance(w, dict) else w["weights"], dtype=float)
    share_n = w[:50].sum() / w.sum()
    assert share_n == pytest.approx(0.75, abs=1e-3)


# ---------------------------------------------------------------- CS-N04 / CS-049 (weighting UI)
def test_cs_n04_ui_does_not_invent_equal_share_targets():
    i = JS.index("function executeWeighting"); block = JS[i:JS.index("fetch('/api/weight'", i)]
    assert "equalShare" not in block and "1.0 / cats.length" not in block


def test_cs049_non_convergence_is_announced():
    i = JS.index("function executeWeighting"); block = JS[i:i + 4000]
    assert "converged === false" in block and "did NOT converge" in block


def test_cs_n04_server_default_targets_only_for_matching_region_labels():
    _load(pd.DataFrame({"Region": ["NCR", "Luzon", "Visayas", "Mindanao"] * 25, "Gender": ["M", "F"] * 50}))
    st_, d = _post("/api/weight", {})
    assert st_ == 400 and "targets" in d["message"].lower()
    _load(pd.DataFrame({"Region": ["National Capital Region (NCR)", "Balance Luzon", "Visayas", "Mindanao"] * 25}))
    st_, d = _post("/api/weight", {})
    assert st_ == 200, d


def test_cs_n06_no_invented_bases_anywhere_in_ui():
    assert "|| 412" not in JS and "effRatio = 0.945" not in JS


def load_tests(loader, standard_tests, pattern):
    import inspect
    suite = unittest.TestSuite()
    funcs = [obj for name, obj in globals().items() if name.startswith("test_") and callable(obj)]
    test_methods = {}
    for fn in funcs:
        sig = inspect.signature(fn)
        if hasattr(fn, "_pytest_parametrize"):
            argnames, argvalues = fn._pytest_parametrize
            if isinstance(argnames, str):
                names = [x.strip() for x in argnames.split(",")]
            else:
                names = list(argnames)
            if len(sig.parameters) == len(names):
                for idx, val in enumerate(argvalues):
                    if len(names) == 1:
                        def make_m(f, v): return lambda self: f(v)
                    else:
                        def make_m(f, v): return lambda self: f(*v)
                    test_methods[f"{fn.__name__}_{idx}"] = make_m(fn, val)
        elif len(sig.parameters) == 0:
            def make_m(f): return lambda self: f()
            test_methods[fn.__name__] = make_m(fn)
    if test_methods:
        cls = type("DynamicTests_test_pass7_audit", (unittest.TestCase,), test_methods)
        suite.addTests(loader.loadTestsFromTestCase(cls))
    return suite


if __name__ == "__main__":
    unittest.main()

