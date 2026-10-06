"""
Automated Regression Test Suite for ClearSight QA Defect Resolutions:
- CS-DEF-01 through CS-DEF-04
- CS-UI-01 through CS-STAT-01
- CS-USER-01 through CS-USER-07
"""
import io
import json
import os
import unittest
import openpyxl
import pandas as pd
import numpy as np
import tempfile

import server
from engine.codeframe_excel_parser import parse_excel_codeframe
from engine.tabulation_engine import build_crosstab_table
from engine.export_engine import generate_thesis_excel_tables
from engine.stats_engine import calculate_chi_square_df, chi_square_independence
from engine.statistical_suite import (
    run_independent_ttest,
    run_paired_ttest,
    run_mann_whitney_u,
    run_wilcoxon_signed_rank,
    run_kruskal_wallis,
    run_correlation_matrix,
    run_chi_square_association,
    run_linear_regression,
    run_ordinal_logistic_regression,
    run_path_analysis_sem,
    run_kruskal_quadrant_analysis
)


class TestDefectResolutions(unittest.TestCase):

    def setUp(self):
        server.load_bundled_sample()

    def test_cs_def_01_codeframe_parser_fstring_syntax(self):
        """CS-DEF-01: Ensure parser parses NET and Subnet headers without f-string backslash errors."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Codes", "Label", "Anchored Verbatims", "DP Instructions"])
        ws.append([None, "GAVE FAVORABLE COMMENTS (NET)", None, None])
        ws.append([None, "Public Transportation (Subnet)", None, None])
        ws.append([101, "MRT trains are cleaner and faster", "Mas mabilis na ang tren ngayon", "Keep"])
        
        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        cf = parse_excel_codeframe(excel_bytes, codeframe_name="transit_study.xlsx")
        self.assertEqual(cf["name"], "transit_study.xlsx")
        self.assertEqual(len(cf["topics"]), 1)
        self.assertEqual(cf["topics"][0]["net"], "GAVE FAVORABLE COMMENTS (NET)")
        self.assertEqual(cf["topics"][0]["codes"]["pos"]["code_id"], 101)

    def test_cs_def_02_export_academic_tables(self):
        """CS-DEF-02: Ensure generate_thesis_excel_tables dispatches and produces a valid openpyxl workbook."""
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        try:
            generate_thesis_excel_tables(tmp_path, "Consumer Perception Analysis")
            self.assertTrue(os.path.exists(tmp_path))
            self.assertGreater(os.path.getsize(tmp_path), 1000)
            wb = openpyxl.load_workbook(tmp_path)
            self.assertIn("Table 4.1 - Demographics", wb.sheetnames)
            self.assertIn("Table 4.2 - CrossTab", wb.sheetnames)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_cs_def_03_analysis_config_dynamic_execution_desync(self):
        """CS-DEF-03: Ensure execute_tabulation dynamically honors ANALYSIS_CONFIG session toggles."""
        handler = server.ClearSightRequestHandler.__new__(server.ClearSightRequestHandler)
        df = server.SESSION["df"]
        self.assertIsNotNone(df)

        with server.SESSION_LOCK:
            server.SESSION["analysis_config"] = {
                "stats": {
                    "rao_scott_2": True,
                    "chi_square": True,
                    "welch_anova": True,
                    "fdr_benjamini_hochberg": True,
                    "fdr_benjamini_yekutieli": False
                },
                "hygiene": {
                    "rim_weighting": True,
                    "quarantine_straightliners": True,
                    "quarantine_speeders": True
                },
                "nlp": {
                    "pii_masking": True,
                    "contrastive_clause_weighting": True,
                    "review_queue_routing": True
                }
            }

        # 1. Chi-square on categorical variable (Age_Generation)
        tab_chi2_on = handler.execute_tabulation(df, ["Total", "Region"], ["Age_Generation"])[0]
        self.assertIsNotNone(tab_chi2_on.get("chi_square"))

        with server.SESSION_LOCK:
            server.SESSION["analysis_config"]["stats"]["chi_square"] = False

        tab_chi2_off = handler.execute_tabulation(df, ["Total", "Region"], ["Age_Generation"])[0]
        self.assertIsNone(tab_chi2_off.get("chi_square"))

        # 2. Rao-Scott RS2 on multi-select checkbox variable (Brand_Preference)
        tab_rs2_on = handler.execute_tabulation(df, ["Total", "Region"], ["Brand_Preference"])[0]
        self.assertIsNotNone(tab_rs2_on.get("mrcv"))

        with server.SESSION_LOCK:
            server.SESSION["analysis_config"]["stats"]["rao_scott_2"] = False

        tab_rs2_off = handler.execute_tabulation(df, ["Total", "Region"], ["Brand_Preference"])[0]
        self.assertIsNone(tab_rs2_off.get("mrcv"))

        # 3. Welch ANOVA on rating scale mean variable (Overall_CSAT)
        tab_anova_on = handler.execute_tabulation(df, ["Total", "Region"], ["Overall_CSAT"], metric="mean")[0]
        self.assertIsNotNone(tab_anova_on.get("anova"))

        with server.SESSION_LOCK:
            server.SESSION["analysis_config"]["stats"]["welch_anova"] = False

        tab_anova_off = handler.execute_tabulation(df, ["Total", "Region"], ["Overall_CSAT"], metric="mean")[0]
        self.assertIsNone(tab_anova_off.get("anova"))

        # 4. RIM Weighting toggle
        with server.SESSION_LOCK:
            server.SESSION["analysis_config"]["hygiene"]["rim_weighting"] = False

        tab_no_rim = handler.execute_tabulation(df, ["Total", "Region"], ["Brand_Preference"])[0]
        self.assertEqual(tab_no_rim["unweighted_bases"][0], tab_no_rim["weighted_bases"][0])

    def test_cs_def_04_toast_svg_icons(self):
        """CS-DEF-04: Ensure app.js contains SVG stroke icons for showToast."""
        app_js_path = os.path.join(os.path.dirname(__file__), "..", "static", "app.js")
        with open(app_js_path, "r", encoding="utf-8") as f:
            content = f.read()

        self.assertIn("svg width=\"16\" height=\"16\"", content)
        self.assertIn("stroke=\"currentColor\"", content)
        self.assertIn("toast-notification", content)

    def test_cs_ui_01_and_05_html_ids(self):
        """CS-UI-01 & CS-USER-05: Ensure no literal \\n or escaped backslashes in HTML attribute IDs."""
        index_path = os.path.join(os.path.dirname(__file__), "..", "static", "index.html")
        with open(index_path, "r", encoding="utf-8") as f:
            content = f.read()
        self.assertNotIn("id=\\\"detected-respondents\\\"", content)
        self.assertIn("id=\"detected-respondents\"", content)
        self.assertIn("table-controls-spacer", content)

    def test_cs_user_01_dynamic_demographics_n(self):
        """CS-USER-01: Verify generate_thesis_excel_tables reflects active dataset N=1000 dynamically."""
        with server.SESSION_LOCK:
            df_1000 = pd.concat([server.SESSION["df"]] * 3, ignore_index=True)[:1000]
            server.SESSION["df"] = df_1000
            server.SESSION["weights"] = None

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name
        try:
            generate_thesis_excel_tables(tmp_path, "Benteng Bigas Dynamic Study")
            wb = openpyxl.load_workbook(tmp_path)
            ws1 = wb["Table 4.1 - Demographics"]
            header_str = ws1.cell(row=3, column=2).value
            self.assertIn("N = 1000", header_str)
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_cs_user_02_high_contrast_table_styles(self):
        """CS-USER-02: Ensure styles.css has #0F172A dark slate headers and zebra striping."""
        styles_path = os.path.join(os.path.dirname(__file__), "..", "static", "styles.css")
        with open(styles_path, "r", encoding="utf-8") as f:
            css = f.read()
        self.assertIn("background-color: #0F172A !important;", css)
        self.assertIn("background-color: #F8FAFC !important;", css)
        self.assertIn("background-color: #F1F5F9 !important;", css)

    def test_cs_user_03_no_fmcg_fallback(self):
        """CS-USER-03: Ensure pick_codeframe returns None rather than silently defaulting to consumer_default."""
        self.assertIsNone(server.pick_codeframe({}))
        self.assertEqual(server.pick_codeframe({"category": "Elections"}), "governance_default")
        self.assertEqual(server.pick_codeframe({"category": "FMCG"}), "consumer_default")

    def test_cs_user_04_custom_netting(self):
        """CS-USER-04: Ensure build_crosstab_table dynamically generates custom Net row."""
        df = server.SESSION["df"]
        custom_nets = [{
            "stub": "Brand_Preference",
            "label": "NET: Top Brands (Brand A + Brand B)",
            "categories": ["Brand A", "Brand B"]
        }]
        t = build_crosstab_table(df, "Brand_Preference", ["Total", "Region"], custom_nets=custom_nets)
        first_row = t["rows"][0]
        self.assertEqual(first_row["label"], "NET: Top Brands (Brand A + Brand B)")
        self.assertTrue(first_row.get("is_net"))

    def test_cs_user_06_codeframe_set_serialization(self):
        """CS-USER-06: Ensure handle_upload_codeframe sanitizes sets so json.dumps succeeds."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Codes", "Label", "Anchored Verbatims", "DP Instructions"])
        ws.append([None, "GAVE FAVORABLE COMMENTS (NET)", None, None])
        ws.append([None, "Quality (Subnet)", None, None])
        ws.append([101, "Great service and speed", "Maganda ang serbisyo", "None"])
        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        handler = server.ClearSightRequestHandler.__new__(server.ClearSightRequestHandler)
        captured = []
        handler.send_json_response = lambda data, status=200: captured.append((status, data))
        handler.headers = {"X-Filename": "custom_dp.xlsx"}
        handler.handle_upload_codeframe(excel_bytes)

        status, resp = captured.pop()
        self.assertEqual(status, 200)
        self.assertEqual(resp.get("status"), "success")

    def test_cs_user_07_quadrant_analysis_autodetect_target(self):
        """CS-USER-07: Ensure handle_stats_quadrant autodetects target when Overall_CSAT is absent."""
        with server.SESSION_LOCK:
            df = server.SESSION["df"].drop(columns=["Overall_CSAT"])
            server.SESSION["df"] = df
            server.SESSION["schema"]["Repurchase_Intent"] = {"type": "rating_scale"}

        handler = server.ClearSightRequestHandler.__new__(server.ClearSightRequestHandler)
        captured = []
        handler.send_json_response = lambda data, status=200: captured.append((status, data))
        handler.handle_stats_quadrant(json.dumps({}).encode("utf-8"))

        status, resp = captured.pop()
        self.assertEqual(status, 200)
        self.assertEqual(resp.get("status"), "success")
        self.assertEqual(resp.get("results", {}).get("target_variable"), "Repurchase_Intent")

    def test_cs_stat_ipa_dynamic_zero_hardcoding(self):
        """CS-STAT-IPA-DYNAMIC: Verify 100% dynamic Kruskal IPA on arbitrary dataset with zero hardcoding."""
        from engine.statistical_suite import run_dynamic_kruskal_quadrant_analysis
        np.random.seed(42)
        n = 120
        df_arb = pd.DataFrame({
            "voting_intent": np.random.choice([1, 2, 3, 4, 5], size=n),
            "traffic_management": np.random.choice([1, 2, 3, 4, 5], size=n),
            "healthcare_access": np.random.choice([1, 2, 3, 4, 5], size=n),
            "anti_corruption": np.random.choice([1, 2, 3, 4, 5], size=n),
            "job_creation": np.random.choice([1, 2, 3, 4, 5], size=n)
        })

        attrs = ["traffic_management", "healthcare_access", "anti_corruption", "job_creation"]
        res = run_dynamic_kruskal_quadrant_analysis(df_arb, "voting_intent", attrs)

        self.assertNotIn("error", res)
        self.assertEqual(res["target_variable"], "voting_intent")
        self.assertEqual(res["sample_size"], n)
        self.assertIn("cutoffs", res)
        self.assertIn("performance_midpoint", res["cutoffs"])
        self.assertEqual(res["cutoffs"]["importance_midpoint"], 25.0)  # 100% / 4 attributes

        for item in res["attributes"]:
            self.assertIn(item["quadrant_code"], ("Q1", "Q2", "Q3", "Q4"))
            self.assertIn(item["quadrant_color"], ("#DC2626", "#16A34A", "#94A3B8", "#F59E0B"))
            self.assertGreater(item["derived_importance_pct"], 0)


    def test_cs_coder_v3_router(self):
        """CS-CODER-V3-ROUTER: Verify Coder v3 environment flag, endpoint routing, and batch execution."""
        # 1. Verify default environment flag
        self.assertEqual(server.CODER_DEFAULT, "v3")

        # 2. Verify /api/code-open-ends dispatches coder v3 successfully
        handler = server.ClearSightRequestHandler.__new__(server.ClearSightRequestHandler)
        captured = []
        handler.send_json_response = lambda data, status=200: captured.append((status, data))
        handler.validate_host_header = lambda: True
        handler.validate_origin_header = lambda: True

        # Simulate POST to /api/code-open-ends without codeframe -> returns needs_codeframe prompt
        handler.path = "/api/code-open-ends"
        handler.headers = {"Content-Length": "15"}
        handler.rfile = io.BytesIO(b'{"coder": "v3"}')
        handler.do_POST()
        self.assertTrue(len(captured) > 0)
        status, resp = captured.pop()
        self.assertEqual(status, 200)
        self.assertEqual(resp.get("status"), "needs_codeframe")

        # Simulate POST to /api/code-open-ends with domain codeframe
        payload = json.dumps({"coder": "v3", "codeframe": "consumer_default"}).encode("utf-8")
        handler.headers = {"Content-Length": str(len(payload))}
        handler.rfile = io.BytesIO(payload)
        handler.do_POST()

        self.assertTrue(len(captured) > 0)
        status, resp = captured.pop()
        self.assertEqual(status, 200)
        self.assertEqual(resp.get("status"), "success")
        self.assertEqual(resp.get("coder"), "v3")
        self.assertEqual(resp.get("coder_version"), "v3")
        self.assertIn("records", resp)
        self.assertIn("sentiment_counts", resp)
        self.assertIn("codeframe", resp)
        self.assertGreater(resp.get("total_analyzed", 0), 0)

        # 3. Verify validation rejects invalid coder like "v4"
        captured.clear()
        handler.rfile = io.BytesIO(b'{"coder": "v4"}')
        handler.headers = {"Content-Length": "15"}
        handler.do_POST()
        self.assertTrue(len(captured) > 0)
        status, resp = captured.pop()
        self.assertEqual(status, 400)
        self.assertEqual(resp.get("status"), "error")
        self.assertIn("coder must be 'v1', 'v2', or 'v3'", resp.get("message", ""))


if __name__ == "__main__":
    unittest.main()
